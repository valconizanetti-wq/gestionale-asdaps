import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
righe = json.load(open(sys.argv[1])); OUT = sys.argv[2]; VERS = sys.argv[3]
F = "Arial"
wb = Workbook()
# --- Istruzioni
ist = wb.active; ist.title = "Istruzioni"
ist.sheet_view.showGridLines = False
ist.column_dimensions["A"].width = 3; ist.column_dimensions["B"].width = 120
testi = [
 ("Testi del Gestionale ASD da correggere", Font(name=F, size=15, bold=True, color="0F4124")),
 (f"Esportazione del codice {VERS}: {len(righe)} testi, presi da pagine, email, PDF, messaggi ed errori.", Font(name=F, size=10, color="555555")),
 ("", None),
 ("Come si usa", Font(name=F, size=11, bold=True)),
 ("1. Nel foglio «Testi» usa i filtri in alto: per Modulo (per esempio «Anticipi e note di rimborso»), per Tipo, o «Citato nel collaudo» (Q, R, S…).", None),
 ("2. Scrivi la versione giusta solo nella colonna gialla «Testo corretto», e solo per i testi da cambiare. Le righe con la colonna gialla vuota restano come sono.", None),
 ("3. Se un testo va ripensato e non sai come scriverlo, lascia «Testo corretto» vuoto e scrivi la richiesta in «Note per Claude» (per esempio «troppo lungo, accorcia»).", None),
 ("4. Rimandami il file: applico le correzioni in tutti i punti indicati in «Dove», aggiorno i collaudi che citano quei testi e le prove automatiche, e preparo la versione nuova.", None),
 ("", None),
 ("Da lasciare come sono, dentro il testo", Font(name=F, size=11, bold=True)),
 ("• le parti variabili: {{ … }}, {$…}, :attribute, :min, {nome} e simili. Puoi spostarle nella frase, ma non cambiarle. Dentro {{ … }} si correggono solo le parole tra apici.", None),
 ("• i tag come <strong>…</strong>, <a …>…</a>, <br>: servono per il grassetto, i link e gli a capo.", None),
 ("• gli spazi all'inizio o alla fine: alcuni testi sono pezzi di una frase che il programma compone (vedi «Dove»).", None),
 ("", None),
 ("Da sapere", Font(name=F, size=11, bold=True)),
 ("• Un testo che compare in più punti dello stesso modulo è su una riga sola: la correzione vale per tutti quei punti (colonna «Punti»). Lo stesso testo in un altro modulo ha la sua riga.", None),
 ("• Alcune righe sono pezzi di frase o righe di un paragrafo lungo: la frase intera si legge nel file indicato in «Dove».", None),
 ("• I testi già salvati nel database (promemoria, categorie, scadenze) non sono qui: si correggono dalle loro pagine del gestionale.", None),
 ("• Il foglio «Posizioni» (nascosto) serve a me per applicare le correzioni: non modificarlo.", None),
 ("", None),
 ("Esempio", Font(name=F, size=11, bold=True)),
 ("Testo attuale: «Copia firmata caricata: ora si registra il pagamento.»  →  Testo corretto: «Copia firmata caricata: ora registra il pagamento.»", None),
]
for i, (t, f) in enumerate(testi, 1):
    c = ist.cell(row=i + 1, column=2, value=t)
    c.font = f or Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
# --- Testi
ws = wb.create_sheet("Testi")
cols = [("ID", 8), ("Modulo", 24), ("Tipo", 20), ("Citato nel collaudo", 12), ("Testo attuale", 70), ("Testo corretto", 70), ("Note per Claude", 30), ("Punti", 7), ("Dove (file:riga)", 48)]
head = Font(name=F, size=10, bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="0F4124")
giallo = PatternFill("solid", fgColor="FFF4C2"); bordo = Border(bottom=Side(style="thin", color="DDDDDD"))
for j, (n, w) in enumerate(cols, 1):
    c = ws.cell(row=1, column=j, value=n); c.font = head; c.fill = fill
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 30
norm = Font(name=F, size=10); wrap = Alignment(wrap_text=True, vertical="top")
for i, r in enumerate(righe, 2):
    vals = [r["id"], " · ".join(r["moduli"]), " · ".join(r["tipi"]), r["collaudo"], r["testo"], None, None, len(r["dove"]), "\n".join(r["dove"][:6]) + (f"\n… e altri {len(r['dove'])-6}" if len(r["dove"]) > 6 else "")]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v); c.font = norm; c.alignment = wrap; c.border = bordo
        if j == 6: c.fill = giallo
    for j in (5, 6):
        ws.cell(row=i, column=j).number_format = "@"
ws.freeze_panes = "F2"
ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(cols)).column_letter}{len(righe) + 1}"
# --- Posizioni (nascosto)
po = wb.create_sheet("Posizioni")
po.append(["ID", "File", "Riga", "Tipo di testo", "Originale nel codice"])
for r in righe:
    for v in r["voci"]:
        po.append([r["id"], v["file"], v["riga"], v["kind"], v["raw"]])
for row in po.iter_rows():
    for c in row: c.font = norm
po.sheet_state = "hidden"
wb.active = 1
wb.save(OUT)
print("ok", len(righe))
